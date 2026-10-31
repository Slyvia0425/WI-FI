"""Read-only TCS-Fall annotation, repository and demonstration-packet audit.

Run with the bundled Python containing openpyxl. No training or upstream code execution.
"""
import argparse
from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import subprocess

import numpy as np
import openpyxl


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def annotations(path):
    before = sha(path)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    if workbook.sheetnames != ['Sheet2']:
        raise ValueError('Workbook structure changed; inspect before adapting')
    sheet = workbook['Sheet2']
    values = list(sheet.values)
    if values[0][:11] != ('ID', 'gender', 'age', 'height', 'weight', 10, 11, 12, 13, 14, 15):
        raise ValueError('Unexpected annotation headers')
    notes = [row[11] for row in values[1:] if len(row) > 11 and row[11]]
    records, errors = [], []
    for row_idx, row in enumerate(values[1:], 2):
        subject = row[0]
        if not isinstance(subject, int):
            errors.append({'row': row_idx, 'reason': 'invalid subject ID'})
            continue
        for col_idx, activity in enumerate(range(10, 16), 6):
            cell = sheet.cell(row_idx, col_idx)
            text = cell.value
            if not isinstance(text, str):
                errors.append({'cell': cell.coordinate, 'reason': 'missing/nontext annotation'})
                continue
            parts = text.split(',')
            if len(parts) != 2:
                errors.append({'cell': cell.coordinate, 'reason': 'expected two repetitions'})
            for repeat, part in enumerate(parts):
                match = re.fullmatch(r'\s*(\d+)\s*/\s*(\d+)\s*', part)
                if not match:
                    errors.append({'cell': cell.coordinate, 'value': part, 'reason': 'unparsed annotation'})
                    continue
                fall, duration = map(int, match.groups())
                if not 0 <= fall <= duration:
                    errors.append({'cell': cell.coordinate, 'reason': 'event outside stated recording duration'})
                records.append({'subject_id': subject, 'activity_id': activity,
                                'repeat_index_in_cell_zero_based': repeat,
                                'recorded_fall_time_seconds': fall, 'recorded_duration_seconds': duration,
                                'sheet': sheet.title, 'cell': cell.coordinate, 'original_cell_value': text,
                                'raw_file_path': None, 'raw_file_match_verified': False,
                                'event_definition': 'recorded fall time; onset/impact meaning not established'})
    workbook.close()
    if sha(path) != before:
        raise ValueError('Source workbook changed')
    keys = [(r['subject_id'], r['activity_id'], r['repeat_index_in_cell_zero_based']) for r in records]
    return records, {
        'source_path': str(path.resolve()), 'source_sha256': before, 'source_unchanged': True,
        'sheet': 'Sheet2', 'header_range': 'A1:K1', 'time_cells': 'F2:K21',
        'source_note': notes, 'subjects': sorted({r['subject_id'] for r in records}),
        'activity_ids': sorted({r['activity_id'] for r in records}), 'annotation_records': len(records),
        'duplicate_subject_activity_repeat_keys': len(keys) - len(set(keys)), 'errors': errors,
        'recorded_duration_range_seconds': [min(r['recorded_duration_seconds'] for r in records),
                                            max(r['recorded_duration_seconds'] for r in records)],
        'recorded_fall_time_range_seconds': [min(r['recorded_fall_time_seconds'] for r in records),
                                             max(r['recorded_fall_time_seconds'] for r in records)],
        'annotation_stated_total_seconds_not_verified_from_raw': sum(r['recorded_duration_seconds'] for r in records),
        'annotation_stated_total_hours_not_verified_from_raw': sum(r['recorded_duration_seconds'] for r in records) / 3600,
        'recorded_event_outside_final_30_seconds_count': sum(r['recorded_fall_time_seconds'] < r['recorded_duration_seconds'] - 30 for r in records),
        'recorded_events_before_120_seconds_count': sum(r['recorded_fall_time_seconds'] < 120 for r in records),
        'integer_second_resolution': True, 'verified_continuous_raw_file_matches': 0,
    }


def decode_packet_csi(payload, ntx, nrx):
    # Intel payload order inspected in pinned official wifilib.py: subcarrier, Rx, Tx.
    output = np.zeros((ntx, nrx, 30), dtype=np.complex64)
    bit = 0
    for sub in range(30):
        bit += 3
        for rx in range(nrx):
            for tx in range(ntx):
                start, shift = divmod(bit, 8)
                word = int.from_bytes(payload[start:start + 3], 'little') >> shift
                real, imag = word & 255, (word >> 8) & 255
                real = real - 256 if real >= 128 else real
                imag = imag - 256 if imag >= 128 else imag
                output[tx, rx, sub] = complex(real, imag)
                bit += 16
    return output


def audit_dat(path):
    data = path.read_bytes()
    pointer, errors, stamps, counts, layouts, decoded = 0, [], [], [], Counter(), []
    codes = Counter()
    while pointer < len(data):
        if pointer + 2 > len(data):
            errors.append({'offset': pointer, 'reason': 'trailing incomplete length header'}); break
        length = int.from_bytes(data[pointer:pointer + 2], 'big')
        if length == 0 or pointer + 2 + length > len(data):
            errors.append({'offset': pointer, 'reason': 'zero/truncated field', 'field_len': length}); break
        frame = data[pointer + 2:pointer + 2 + length]
        pointer += 2 + length
        codes[frame[0]] += 1
        if frame[0] != 187:
            continue
        if len(frame) < 21:
            errors.append({'offset': pointer - length - 2, 'reason': 'short CSI header'}); continue
        ntx, nrx = frame[10], frame[9]
        expected = (30 * (ntx * nrx * 16 + 3) + 6) // 8
        reported = int.from_bytes(frame[17:19], 'little')
        if not (1 <= ntx <= 3 and 1 <= nrx <= 3 and reported == expected and len(frame[21:]) == expected):
            errors.append({'offset': pointer - length - 2, 'reason': 'invalid CSI dimensions or payload size'}); continue
        layouts[f'{ntx}x{nrx}x30'] += 1
        stamps.append(int.from_bytes(frame[1:5], 'little'))
        counts.append(int.from_bytes(frame[5:7], 'little'))
        if len(decoded) < 128:
            decoded.append(decode_packet_csi(frame[21:], ntx, nrx))
    stamps = np.array(stamps, dtype=np.int64)
    diffs_raw = np.diff(stamps)
    diffs = np.where(diffs_raw < -(2**31), diffs_raw + 2**32, diffs_raw)
    count_diff = (np.diff(np.array(counts, dtype=np.int64)) + 2**16) % 2**16
    amplitude = np.abs(np.stack(decoded)) if decoded else None
    return {
        'file': str(path.resolve()), 'bytes': len(data), 'sha256': sha(path),
        'git_blob_sha1': hashlib.sha1(b'blob ' + str(len(data)).encode() + b'\0' + data).hexdigest(),
        'field_codes_and_counts': dict(codes), 'valid_csi_packet_count': len(stamps),
        'packet_layout_tx_rx_subcarrier_counts': dict(layouts), 'errors': errors,
        'timestamp_units_source': 'https://dhalperi.github.io/linux-80211n-csitool/faq.html#timestamp_low: low 32 bits of 1 MHz NIC clock',
        'timestamp_low_first': int(stamps[0]), 'timestamp_low_last': int(stamps[-1]),
        'timestamp_span_seconds': float(diffs.sum() / 1e6),
        'median_packet_interval_microseconds': float(np.median(diffs)),
        'p99_packet_interval_microseconds': float(np.percentile(diffs, 99)),
        'observed_packets_per_second_over_timestamp_span': float((len(stamps) - 1) / (diffs.sum() / 1e6)),
        'timestamp_rollover_count': int((diffs_raw < -(2**31)).sum()),
        'nonincreasing_timestamp_count_after_rollover_handling': int((diffs <= 0).sum()),
        'bfee_count_nonunit_steps': int((count_diff != 1).sum()),
        'bfee_count_missing_measurements_inferred_if_forward_progress': int(np.maximum(count_diff - 1, 0).sum()),
        'decoded_packet_subset_count': len(decoded), 'decoded_subset_dtype': 'complex64',
        'decoded_subset_raw_amplitude_finite_ratio': float(np.isfinite(amplitude).mean()),
        'decoded_subset_raw_amplitude_range': [float(amplitude.min()), float(amplitude.max())],
        'decode_scope': 'first 128 valid packets; raw unscaled CSI, no antenna permutation/scaling, no model input prepared',
        'continuous_recording': False,
        'label_and_subject_from_filename': 'not asserted; preprocessing notebook explaining naming is missing',
    }


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--source-root', type=Path, required=True)
    ap.add_argument('--git-root', type=Path, required=True)
    ap.add_argument('--output', type=Path, required=True)
    args = ap.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    records, annotation_summary = annotations(args.source_root / 'actual fall time.xlsx')
    tree_text = subprocess.check_output(['git', '-C', str(args.git_root), 'ls-tree', '-r', '-l', 'HEAD'], text=True)
    tree = []
    for line in tree_text.splitlines():
        prefix, path = line.split('\t', 1)
        mode, kind, blob, size = prefix.split()
        tree.append({'path': path, 'git_blob_sha1': blob, 'bytes': int(size)})
    demo = [audit_dat(p) for p in sorted((args.source_root / 'dataset').rglob('*.dat'))]
    for row in demo:
        name = str(Path(row['file']).relative_to(args.source_root.resolve()))
        expected = next(item for item in tree if item['path'] == name)
        if row['git_blob_sha1'] != expected['git_blob_sha1']:
            raise ValueError('Downloaded demo does not match pinned official Git blob')
    demo_content_counts = Counter(r['sha256'] for r in demo)
    summary = {
        'completed_utc': datetime.now(timezone.utc).isoformat(),
        'repository': 'https://github.com/zhouziyu2050/TCS-Fall',
        'commit': subprocess.check_output(['git', '-C', str(args.git_root), 'rev-parse', 'HEAD'], text=True).strip(),
        'repository_tree': tree, 'repository_dat_file_count': sum(r['path'].endswith('.dat') for r in tree),
        'preprocess_notebook_present_in_pinned_repository': any(Path(r['path']).name == 'preprocess.ipynb' for r in tree),
        'annotations': annotation_summary, 'downloaded_demonstration_audits': demo,
        'exact_duplicate_demonstration_file_groups': [digest for digest, n in demo_content_counts.items() if n > 1],
        'continuous_raw_recordings_downloaded': 0, 'model_training_run': False,
        'event_recall': None, 'false_alarms_per_hour': None, 'alert_latency': None,
        'access_status': {
            'ieee_dataport': 'https://ieee-dataport.org/documents/tcs-fall',
            'ieee_observed': 'LOGIN TO ACCESS DATASET FILES; four listed archives and annotation workbook',
            'baidu': 'https://pan.baidu.com/s/1IT5KZEhvhNjuBjQ3EKxMnw?pwd=qitg',
            'baidu_observed': 'dataset directory lists 1.zip to 20.zip; download click opened BaiduNetdisk Mac installer URL, not a raw-data archive',
            'account_login_or_local_archives_required': True,
        },
        'remaining_requirements': [
            'Obtain raw continuous .dat files; match each to subject/activity/repetition without guessing filenames.',
            'Obtain/inspect preprocessing notebook named in official README but absent from pinned public Git tree.',
            'Verify event time reference and onset/impact semantics; integer-second labels do not support millisecond event-latency accuracy.',
            'Verify raw duration, timestamps, packet loss, radio configuration and annotation alignment for each continuous recording.',
            'Keep Intel5300 native features separate from HP232; no reuse of HP normalization by zero padding.',
            'Use independent subject/session partitions before creating windows; tune alarms on validation only.',
        ],
    }
    (args.output / 'annotations_extracted.json').write_text(json.dumps(records, indent=2))
    (args.output / 'audit.json').write_text(json.dumps(summary, indent=2, ensure_ascii=False))
    print(json.dumps({'annotations': annotation_summary,
                      'repository_dat_file_count': summary['repository_dat_file_count'],
                      'demonstration_audits': demo,
                      'continuous_raw_recordings_downloaded': 0}, indent=2, ensure_ascii=False))


if __name__ == '__main__':
    main()
